#!/usr/bin/env python3
"""Temporary: inspect candidate source files (run in GitHub Actions; removed before merge)."""
import io
import re
import sys
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
import openpyxl

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
      "Accept-Language": "ja,en;q=0.8"}


def get(url):
    r = requests.get(url, headers=UA, timeout=60)
    print("GET", url, "->", r.status_code, r.headers.get("content-type"), len(r.content))
    return r


def links(url, pat=None, text_pat=None, limit=60):
    r = get(url)
    soup = BeautifulSoup(r.content, "html.parser")
    out = []
    for a in soup.find_all("a", href=True):
        href = urljoin(r.url, a["href"])
        text = re.sub(r"\s+", " ", a.get_text(" ", strip=True))[:140]
        if (pat and re.search(pat, href)) or (text_pat and re.search(text_pat, text)):
            out.append((text, href))
    for t, h in out[:limit]:
        print("  LINK", t, "|", h)
    return out


def xlsx_summary(url, rows=25):
    r = get(url)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        print("  not xlsx:", exc, r.content[:200]); return
    for ws in wb.worksheets[:4]:
        print("  SHEET", ws.title, ws.max_row, ws.max_column)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= rows:
                break
            vals = [v for v in row if v is not None]
            if vals:
                print("   ", i, vals[:14])


def section(title):
    print("=" * 30, title, "=" * 30)


section("JNTO")
for u in ["https://www.jnto.go.jp/statistics/data/visitors-statistics/"]:
    found = links(u, pat=r"\.(xlsx?|csv|pdf)", text_pat=r"国籍|年表|訪日外客数")
    xs = [h for t, h in found if re.search(r"\.xlsx?$", h)]
    for h in xs[:2]:
        xlsx_summary(h, rows=15)

section("ISA zairyu (latest lids)")
found = links("https://www.moj.go.jp/isa/policies/statistics/toukei_ichiran_touroku.html", pat=r"lid=", limit=0)
print("  N", len(found), "LAST", found[-6:])
for t, h in found[-2:]:
    section("e-Stat datalist " + t)
    fl = links(h, pat=r"file-download|stat-search/files\?.*lid", text_pat=r"国籍", limit=40)

section("ISA nyukan 2025")
fl = links("https://www.e-stat.go.jp/stat-search/files?page=1&layout=datalist&lid=000001486941", pat=r"file-download", text_pat=r"国籍", limit=40)

section("JETRO FDI")
xlsx_summary("https://www.jetro.go.jp/ext_images/world/japan/stats/fdi/data/country1_25cy.xlsx", rows=60)
sys.exit(0)
