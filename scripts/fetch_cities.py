#!/usr/bin/env python3
"""Collect major-city data (regional GDP per capita and city population) with provenance.

Cities and the regional unit each one maps to are listed in basic/cities/cities.json.
Sources (registered in basic/sources.json):
  psa_ph_pcgdp       PSA OpenSTAT: per capita GDP by province and HUC (current prices, PHP)
  nso_th_gpp         NSO Thailand Statistical Yearbook, Table 10.14: GPP per capita (current, THB)
  dosm_state_gdp     DOSM / data.gov.my: state GDP at constant 2015 prices and state population
  un_wup2025_cities  UN World Urbanization Prospects 2025, File 21: population of cities
  esri_kenmin        Cabinet Office (ESRI) prefectural accounts: nominal GDP and population of
                     the designated cities (政令指定都市) in Japan, fiscal years
  worldbank_wdi      exchange rates (PA.NUS.FCRF) and national GDP per capita (current LCU / USD)

Writes:
  basic/cities/raw/<source_id>.json   the rows used, with URL and retrieval time
                                      (left untouched when only retrieved_at would change)
  basic/data/city_observations.{json,csv}
      one row per city, indicator and year. Indicators:
        grdp_per_capita_local  regional GDP per capita in local currency (as published)
        grdp_per_capita_usd    the same converted with the World Bank official exchange rate
        grdp_per_capita_ratio  ratio to the national average from the same basis
        city_population        UN WUP 2025 city population (2025 onward: projections)
      Every derived value says how it was derived (`derivation`).
  basic/cities/fetch_report.json      per-source status and UN name-matching diagnostics

Usage:
  python scripts/fetch_cities.py            # fetch + build
  python scripts/fetch_cities.py --build    # rebuild from saved raw files only
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import time
from datetime import datetime, timezone

import requests

from fetch_basic import BASIC_DIR, load_json, write_json

RAW = BASIC_DIR / "cities" / "raw"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
    "Accept": "*/*",
}
UN_URL = "https://population.un.org/wup/assets/Download/Cities/WUP2025-F21-DEGURBA-Cities_Pop.xlsx"
PSA_URL = "https://openstat.psa.gov.ph/PXWeb/api/v1/en/DB/2A/PPA/0092A5FPPA8.px"
TH_URL = ("https://www.nso.go.th/public/e-book/Statistical-Yearbook/SYB-{y}_webPage/"
          "File%20SYB/{y}/10/T%2010.14%20{y}.pdf")
DOSM_URL = "https://api.data.gov.my/data-catalogue"
ESRI_PAGE = "https://www.esri.cao.go.jp/jp/sna/data/data_list/kenmin/files/contents/main_{y}.html"
ESRI_TABLE = "https://www.esri.cao.go.jp/jp/sna/data/data_list/kenmin/files/contents/tables/{y}/{name}.xlsx"
WB_URL = "https://api.worldbank.org/v2/country/{iso3}/indicator/{code}?format=json&per_page=20000&date=2000:{y}"
FIRST_YEAR = 2000


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def get_json(session: requests.Session, url: str, **kw):
    """JSON from a response that may start with a UTF-8 BOM (PSA OpenSTAT does)."""
    return json.loads(get(session, url, **kw).content.decode("utf-8-sig"))


def get(session: requests.Session, url: str, method: str = "GET", **kw) -> requests.Response:
    for attempt in range(3):
        try:
            resp = session.request(method, url, headers=HEADERS, timeout=120, **kw)
            resp.raise_for_status()
            return resp
        except (requests.Timeout, requests.ConnectionError):
            if attempt == 2:
                raise
            time.sleep(5 * (attempt + 1))
    raise RuntimeError("unreachable")


def is_empty(rows) -> bool:
    """No rows at all, or (for sources made of several series) any series empty."""
    if isinstance(rows, dict):
        return not rows or any(not v for v in rows.values())
    return not rows


def save_raw(source_id: str, new: dict) -> dict:
    """Write raw/<source_id>.json unless only retrieved_at changed.

    An empty result is treated as a failed fetch: the saved data is kept and
    the error is raised so the run reports it instead of committing deletions."""
    path = RAW / f"{source_id}.json"
    old = load_json(path) if path.exists() else None
    if is_empty(new.get("rows")):
        raise ValueError("empty result; kept the previously saved data")
    strip = lambda d: {k: v for k, v in d.items() if k != "retrieved_at"}  # noqa: E731
    if old and strip(old) == strip(new):
        return old
    write_json(path, new)
    return new


# --------------------------------------------------------------------------- #
# Fetchers: each returns {source_id, url(s), retrieved_at, rows, notes}
# --------------------------------------------------------------------------- #

def fetch_psa(session, cities) -> dict:
    meta = get_json(session, PSA_URL)
    geo = next(v for v in meta["variables"] if v["code"] == "Geolocation")
    val = next(v for v in meta["variables"] if v["code"] == "Type of Valuation")
    year = next(v for v in meta["variables"] if v["code"] == "Year")
    wanted = {c["region"]["unit"] for c in cities if c.get("region") and c["region"]["source_id"] == "psa_ph_pcgdp"}
    geo_codes = {t: code for code, t in zip(geo["values"], geo["valueTexts"]) if t in wanted}
    current = next(code for code, t in zip(val["values"], val["valueTexts"]) if t == "At Current Prices")
    query = {"query": [
        {"code": "Geolocation", "selection": {"filter": "item", "values": list(geo_codes.values())}},
        {"code": "Type of Valuation", "selection": {"filter": "item", "values": [current]}},
        {"code": "Year", "selection": {"filter": "item", "values": year["values"]}},
    ], "response": {"format": "json"}}
    data = get_json(session, PSA_URL, method="POST", json=query)
    code_to_geo = {v: k for k, v in geo_codes.items()}
    code_to_year = dict(zip(year["values"], year["valueTexts"]))
    rows = []
    for d in data.get("data", []):
        g, _, y = d["key"]
        v = d["values"][0]
        if v not in ("", ".", "..", None):
            rows.append({"unit": code_to_geo[g], "year": int(code_to_year[y]), "value": float(v)})
    return {"source_id": "psa_ph_pcgdp", "url": PSA_URL, "retrieved_at": now(), "currency": "PHP",
            "basis": "current", "table": meta.get("title"), "rows": rows,
            "missing_units": sorted(wanted - set(geo_codes))}


def fetch_th(session, cities) -> dict:
    import pypdf
    wanted = {c["region"]["unit"] for c in cities if c.get("region") and c["region"]["source_id"] == "nso_th_gpp"}
    errors = []
    for y in range(datetime.now(timezone.utc).year + 1, 2022, -1):
        url = TH_URL.format(y=y)
        try:
            body = get(session, url).content
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{y}: {exc}")
            continue
        reader = pypdf.PdfReader(io.BytesIO(body))
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
        years = [int(x) for x in re.findall(r"\((\d{4})\)", text)[:3]]
        rows, national = [], None
        line_re = re.compile(r"((?:[\d,]+\s+){5}[\d,]+)\s*([A-Za-z][A-Za-z .\-]*)?\s*$")
        for line in text.splitlines():
            m = line_re.search(line)
            if not m:
                continue
            nums = [float(n.replace(",", "")) for n in m.group(1).split()]
            name = (m.group(2) or "").strip()
            per_capita = nums[3:]
            if not name and national is None:
                national = per_capita  # first unnamed row = whole kingdom
            elif name in wanted:
                rows += [{"unit": name, "year": yr, "value": v} for yr, v in zip(years, per_capita)]
        if national:
            rows += [{"unit": "Whole Kingdom", "year": yr, "value": v} for yr, v in zip(years, national)]
        return {"source_id": "nso_th_gpp", "url": url, "retrieved_at": now(), "currency": "THB", "basis": "current",
                "table": "Table 10.14 Gross Regional and Provincial Product by Current Market Prices",
                "rows": rows, "missing_units": sorted(wanted - {r["unit"] for r in rows}),
                "note": "最新年は暫定値（p） / latest year is preliminary"}
    raise RuntimeError("no yearbook found: " + "; ".join(errors))


def fetch_dosm(session, cities) -> dict:
    def catalogue(dataset: str, **params):
        resp = get(session, DOSM_URL, params={"id": dataset, "limit": 100000, **params})
        return resp.json(), resp.url
    gdp, gdp_url = catalogue("gdp_state_real_supply")
    pop, pop_url = catalogue("population_state")
    gdp_rows = [{"state": r["state"], "year": int(r["date"][:4]), "value": r["value"]}
                for r in gdp if r.get("sector") == "p0" and r.get("series") == "abs"]
    pop_rows = [{"state": r["state"], "year": int(r["date"][:4]), "value": r["population"]}
                for r in pop if r.get("age") == "overall_age" and r.get("sex") == "overall_sex"
                and r.get("ethnicity") == "overall_ethnicity"]
    return {"source_id": "dosm_state_gdp", "urls": {"gdp": gdp_url, "population": pop_url}, "url": gdp_url,
            "retrieved_at": now(), "currency": "MYR", "basis": "constant_2015",
            "rows": {"gdp_rm_million": gdp_rows, "population_thousand": pop_rows},
            "states": sorted({r["state"] for r in gdp_rows})}


def fetch_un(session, cities, iso3s) -> dict:
    import openpyxl
    body = get(session, UN_URL).content
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True)
    ws = wb["Data"]
    header = None
    rows, by_country = [], {}
    for r in ws.iter_rows(values_only=True):
        if header is None:
            if r and r[0] == "Index":
                header = list(r)
            continue
        rec = dict(zip(header, r))
        if rec.get("ISO3_Code") not in iso3s:
            continue
        years = {int(k): v for k, v in rec.items() if isinstance(k, (int, str)) and str(k).isdigit()
                 and int(k) >= FIRST_YEAR and v is not None}
        by_country.setdefault(rec["ISO3_Code"], []).append((rec["City_Name"], years.get(2025) or 0))
        rows.append({"iso3": rec["ISO3_Code"], "city_code": rec.get("City_Code"), "city_name": rec["City_Name"],
                     "capital": rec.get("Capital"), "population_thousand": {str(k): v for k, v in years.items()}})
    names = {c["id"]: [n.lower() for n in c["un_names"]] for c in cities}
    # Keep the named cities plus every city of 300,000+ (2025), so a name fix needs no re-fetch.
    kept = [r for r in rows if any(r["city_name"].lower() in ns for ns in names.values())
            or (r["population_thousand"].get("2025") or 0) >= 300]
    largest = {k: [n for n, _ in sorted(v, key=lambda t: -t[1])[:12]] for k, v in by_country.items()}
    return {"source_id": "un_wup2025_cities", "url": UN_URL, "retrieved_at": now(), "rows": kept,
            "largest_cities_by_country": largest}


def fetch_esri(session, cities) -> dict:
    """Nominal GDP (million yen) and total population of designated cities, from the latest release."""
    import openpyxl
    squash = lambda t: re.sub(r"\s+", "", t)  # noqa: E731 - labels may be spaced out ("札　幌　市")
    wanted = {c["region"]["unit"] for c in cities if c.get("region") and c["region"]["source_id"] == "esri_kenmin"}
    for y in range(datetime.now(timezone.utc).year, 2019, -1):
        try:
            get(session, ESRI_PAGE.format(y=y))
        except Exception:  # noqa: BLE001 - release not published yet
            continue
        tables = {}
        for name in ("soukatu1", "soukatu9"):  # GDP (production, nominal); total population
            url = ESRI_TABLE.format(y=y, name=name)
            ws = openpyxl.load_workbook(io.BytesIO(get(session, url).content), read_only=True, data_only=True).worksheets[0]
            years, rows, seen = None, [], []
            for r in ws.iter_rows(values_only=True):
                if years is None:
                    if r and isinstance(r[3], (int, float)) and 1990 < r[3] < 2100:
                        years = [int(v) for v in r[3:] if isinstance(v, (int, float))]
                    continue
                label = squash(r[1]) if isinstance(r[1], str) else ""
                if label:
                    seen.append(label)
                if label in wanted or label == "全県計":
                    rows += [{"unit": label, "year": yr, "value": v}
                             for yr, v in zip(years, r[3:3 + len(years)]) if isinstance(v, (int, float))]
            tables[name] = {"url": url, "rows": rows, "seen": seen}
        return {"source_id": "esri_kenmin", "url": ESRI_PAGE.format(y=y), "retrieved_at": now(), "currency": "JPY",
                "basis": "current", "release": f"{y}年度版", "urls": {k: v["url"] for k, v in tables.items()},
                "rows": {"gdp_million_yen": tables["soukatu1"]["rows"], "population": tables["soukatu9"]["rows"]},
                "missing_units": sorted(wanted - {r["unit"] for r in tables["soukatu1"]["rows"]}),
                "labels_seen": tables["soukatu1"]["seen"][-40:]}
    raise RuntimeError("no ESRI prefectural accounts release found")


def fetch_wb(session, iso3s) -> dict:
    y = datetime.now(timezone.utc).year
    rows = {}
    for code in ("PA.NUS.FCRF", "NY.GDP.PCAP.CN", "NY.GDP.PCAP.CD"):
        url = WB_URL.format(iso3=";".join(iso3s), code=code, y=y)
        _, data = get(session, url).json()
        rows[code] = [{"iso3": d["countryiso3code"], "year": int(d["date"]), "value": d["value"]}
                      for d in data or [] if d["value"] is not None]
        time.sleep(1)
    return {"source_id": "worldbank_wdi", "url": WB_URL.format(iso3=";".join(iso3s), code="{indicator}", y=y),
            "retrieved_at": now(), "rows": rows}


# --------------------------------------------------------------------------- #
# Build
# --------------------------------------------------------------------------- #

def build(cities, countries, sources, raw) -> list[dict]:
    by_code = {c["code"]: c for c in countries}
    wb = raw.get("worldbank_wdi", {}).get("rows", {})

    def wb_value(code, iso3, year):
        return next((r["value"] for r in wb.get(code, []) if r["iso3"] == iso3 and r["year"] == year), None)

    def wb_url(code):
        return raw["worldbank_wdi"]["url"].replace("{indicator}", code) if "worldbank_wdi" in raw else None

    out = []

    def row(city, indicator, unit, year, value, source_id, url, retrieved_at, **extra):
        c = by_code[city["country"]]
        s = sources.get(source_id, {})
        labels = {"grdp_per_capita_local": ("一人当たり地域GDP（現地通貨）", "Regional GDP per capita (local currency)"),
                  "grdp_per_capita_usd": ("一人当たり地域GDP（米ドル換算）", "Regional GDP per capita (USD)"),
                  "grdp_per_capita_ratio": ("一人当たり地域GDPの対全国比", "Regional GDP per capita, ratio to national"),
                  "city_population": ("都市人口", "City population")}[indicator]
        out.append({
            "city_id": city["id"], "city_ja": city["name_ja"], "city_en": city["name_en"],
            "country_code": city["country"], "country_ja": c["name_ja"], "country_en": c["name_en"],
            "indicator_id": indicator, "indicator_ja": labels[0], "indicator_en": labels[1], "unit": unit,
            "year": year, "value": value,
            "region_unit": (city.get("region") or {}).get("unit") if indicator != "city_population" else extra.pop("region_unit", None),
            "unit_note": (city.get("region") or {}).get("unit_note") if indicator != "city_population" else extra.pop("unit_note", None),
            "note": extra.pop("note", None), "derivation": extra.pop("derivation", None),
            "source_id": source_id, "source_publisher_ja": s.get("publisher_ja"), "source_publisher_en": s.get("publisher_en"),
            "source_title_ja": s.get("title_ja"), "license": s.get("license"),
            "url": url, "retrieved_at": retrieved_at, "retrieval_method": "auto", **extra,
        })

    def add_regional(city, series, national, src, currency, basis, nat_label, extra_note=None, local_derivation=None):
        """series/national: {year: per-capita value in local currency}."""
        iso3 = by_code[city["country"]]["iso3"]
        basis_note = "名目（当年価格）" if basis == "current" else "実質（2015年価格）"
        if extra_note:
            basis_note += "・" + extra_note
        for year, v in sorted(series.items()):
            row(city, "grdp_per_capita_local", currency, year, round(v, 2), src["source_id"], src["url"],
                src["retrieved_at"], note=basis_note,
                derivation=local_derivation or (None if basis == "current" else "州の実質GDP（2015年価格）÷ 州の人口（同じ統計局のデータ）"))
            fx = wb_value("PA.NUS.FCRF", iso3, year) if basis == "current" else None
            if fx:
                row(city, "grdp_per_capita_usd", "USD", year, round(v / fx, 2), src["source_id"], src["url"],
                    src["retrieved_at"], note=basis_note,
                    derivation=f"{v:,.0f} {currency} ÷ 世界銀行 公式為替レート PA.NUS.FCRF {year}年平均 {fx:,.4f} {currency}/USD",
                    fx_rate=fx, fx_url=wb_url("PA.NUS.FCRF"))
            nat = national.get(year)
            if nat:
                row(city, "grdp_per_capita_ratio", "ratio", year, round(v / nat, 3), src["source_id"], src["url"],
                    src["retrieved_at"], note=basis_note,
                    derivation=f"{v:,.0f} ÷ 全国 {nat:,.0f} {currency}（{nat_label}）")

    # Philippines
    if "psa_ph_pcgdp" in raw:
        src = raw["psa_ph_pcgdp"]
        national = {r["year"]: r["value"] for r in wb.get("NY.GDP.PCAP.CN", []) if r["iso3"] == "PHL"}
        for city in cities:
            if (city.get("region") or {}).get("source_id") == "psa_ph_pcgdp":
                series = {r["year"]: r["value"] for r in src["rows"] if r["unit"] == city["region"]["unit"]}
                add_regional(city, series, national, src, "PHP", "current", "世界銀行 NY.GDP.PCAP.CN")
    # Thailand
    if "nso_th_gpp" in raw:
        src = raw["nso_th_gpp"]
        national = {r["year"]: r["value"] for r in src["rows"] if r["unit"] == "Whole Kingdom"}
        for city in cities:
            if (city.get("region") or {}).get("source_id") == "nso_th_gpp":
                series = {r["year"]: r["value"] for r in src["rows"] if r["unit"] == city["region"]["unit"]}
                add_regional(city, series, national, src, "THB", "current", "同じ表の全国値")
    # Malaysia: constant 2015 prices, GDP (RM million) / population (thousand) -> RM per person
    if "dosm_state_gdp" in raw:
        src = raw["dosm_state_gdp"]
        gdp, pop = src["rows"]["gdp_rm_million"], src["rows"]["population_thousand"]
        years = sorted({r["year"] for r in gdp} & {r["year"] for r in pop})
        national = {}
        for y in years:
            g = sum(r["value"] for r in gdp if r["year"] == y)  # includes the extra-regional (Supra) row
            p = sum(r["value"] for r in pop if r["year"] == y and r["state"].lower() != "malaysia")
            national[y] = g / p * 1000 if p else None
        for city in cities:
            if (city.get("region") or {}).get("source_id") != "dosm_state_gdp":
                continue
            unit = city["region"]["unit"]
            # GDP for Kuala Lumpur includes Putrajaya, so divide by both populations.
            pop_states = {unit, "W.P. Putrajaya"} if "Kuala Lumpur" in unit else {unit}
            series = {}
            for y in years:
                g = next((r["value"] for r in gdp if r["year"] == y and r["state"] == unit), None)
                p = sum(r["value"] for r in pop if r["year"] == y and r["state"] in pop_states)
                if g is not None and p:
                    series[y] = g / p * 1000
            add_regional(city, series, national, {**src, "url": src["urls"]["gdp"]}, "MYR", "constant_2015",
                         "全州の合計 ÷ 全州人口")
    # Japan: designated cities, fiscal years; GDP (million yen) / population
    if "esri_kenmin" in raw:
        src = raw["esri_kenmin"]
        gdp, pop = src["rows"]["gdp_million_yen"], src["rows"]["population"]

        def per_capita(unit):
            p = {r["year"]: r["value"] for r in pop if r["unit"] == unit}
            return {r["year"]: r["value"] * 1e6 / p[r["year"]] for r in gdp if r["unit"] == unit and p.get(r["year"])}
        national = per_capita("全県計")
        for city in cities:
            if (city.get("region") or {}).get("source_id") == "esri_kenmin":
                add_regional(city, per_capita(city["region"]["unit"]), national, src, "JPY", "current",
                             "同じ表の全県計 ÷ 全県人口", extra_note="年度（4月〜翌3月）の値を暦年平均レートで換算",
                             local_derivation="市内総生産（名目）÷ 市の総人口（内閣府「県民経済計算」総括表1・9）")
    # Singapore: the city is the country
    for city in cities:
        if (city.get("region") or {}).get("source_id") == "national" and wb:
            iso3 = by_code[city["country"]]["iso3"]
            for r in wb.get("NY.GDP.PCAP.CD", []):
                if r["iso3"] == iso3:
                    row(city, "grdp_per_capita_usd", "USD", r["year"], r["value"], "worldbank_wdi",
                        wb_url("NY.GDP.PCAP.CD"), raw["worldbank_wdi"]["retrieved_at"], note="都市国家のため国全体の値")
                    row(city, "grdp_per_capita_ratio", "ratio", r["year"], 1.0, "worldbank_wdi",
                        wb_url("NY.GDP.PCAP.CD"), raw["worldbank_wdi"]["retrieved_at"], note="都市国家のため国全体と同じ")
    # City population (UN WUP 2025)
    if "un_wup2025_cities" in raw:
        src = raw["un_wup2025_cities"]
        for city in cities:
            names = [n.lower() for n in city["un_names"]]
            iso3 = by_code[city["country"]]["iso3"]
            match = next((r for n in names for r in src["rows"] if r["iso3"] == iso3 and r["city_name"].lower() == n), None)
            if not match:
                continue
            for y, v in sorted(match["population_thousand"].items(), key=lambda kv: int(kv[0])):
                year = int(y)
                row(city, "city_population", "persons", year, round(v * 1000), "un_wup2025_cities", src["url"],
                    src["retrieved_at"], region_unit=match["city_name"],
                    unit_note=f"国連 WUP 2025 の都市（DEGURBA の都市域）: {match['city_name']}（City_Code {match['city_code']}）",
                    note="2025年以降は国連の予測値" if year > 2025 else None)
    return out


def write_table(name: str, rows: list[dict]) -> None:
    out = BASIC_DIR / "data"
    (out / f"{name}.json").write_text(json.dumps(rows, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    columns = []
    for r in rows:
        columns += [k for k in r if k not in columns]
    with open(out / f"{name}.csv", "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--build", action="store_true", help="rebuild from saved raw files only")
    args = parser.parse_args()

    cities = load_json(BASIC_DIR / "cities" / "cities.json")["cities"]
    master = load_json(BASIC_DIR / "countries.json")
    # Benchmarks (Japan) are included for the designated-city comparison.
    countries = master["countries"] + master.get("benchmarks", [])
    sources = {s["id"]: s for s in load_json(BASIC_DIR / "sources.json")["sources"]}
    iso3s = [c["iso3"] for c in countries]

    report = {}
    if not args.build:
        session = requests.Session()
        jobs = {
            "psa_ph_pcgdp": lambda: fetch_psa(session, cities),
            "nso_th_gpp": lambda: fetch_th(session, cities),
            "dosm_state_gdp": lambda: fetch_dosm(session, cities),
            "un_wup2025_cities": lambda: fetch_un(session, cities, iso3s),
            "esri_kenmin": lambda: fetch_esri(session, cities),
            "worldbank_wdi": lambda: fetch_wb(session, ["MYS", "THA", "PHL", "SGP", "VNM", "IDN", "JPN"]),
        }
        for source_id, job in jobs.items():
            try:
                saved = save_raw(source_id, job())
                report[source_id] = {"fetch_status": "ok", "rows": len(saved["rows"]) if isinstance(saved["rows"], list)
                                     else {k: len(v) for k, v in saved["rows"].items()}}
                for key in ("missing_units", "states", "largest_cities_by_country", "release", "labels_seen"):
                    if key in saved:
                        report[source_id][key] = saved[key]
            except Exception as exc:  # noqa: BLE001
                report[source_id] = {"fetch_status": "error", "fetch_error": f"{type(exc).__name__}: {exc}"}
            print(f"{source_id:<20} {report[source_id].get('fetch_status')} {report[source_id].get('fetch_error', '')}")

    raw = {p.stem: load_json(p) for p in sorted(RAW.glob("*.json"))}
    rows = build(cities, countries, sources, raw)
    write_table("city_observations", rows)
    by_code = {c["code"]: c for c in countries}
    write_table("city_list", [{
        "city_id": c["id"], "city_ja": c["name_ja"], "city_en": c["name_en"], "country_code": c["country"],
        "country_ja": by_code[c["country"]]["name_ja"], "designated_city": c.get("designated_city", False),
        "region_unit": (c.get("region") or {}).get("unit"), "unit_note": (c.get("region") or {}).get("unit_note"),
        "gdp_note": c.get("gdp_note"), "un_note": c.get("un_note"),
    } for c in cities])
    covered = {}
    for r in rows:
        covered.setdefault(r["city_id"], set()).add(r["indicator_id"])
    report["coverage"] = {c["id"]: sorted(covered.get(c["id"], [])) for c in cities}
    if not args.build:
        write_json(BASIC_DIR / "cities" / "fetch_report.json", report)
    for c in cities:
        print(f"{c['id']:<7} {', '.join(sorted(covered.get(c['id'], []))) or '(no data)'}")
    return 0 if args.build or any(v.get("fetch_status") == "ok" for k, v in report.items() if k != "coverage") else 1


if __name__ == "__main__":
    sys.exit(main())
