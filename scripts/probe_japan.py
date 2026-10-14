#!/usr/bin/env python3
"""Temporary probe for Japanese city data sources (removed before merge)."""
import io
from pathlib import Path
import requests
import openpyxl

OUT = Path(__file__).resolve().parent.parent / "basic" / "cities" / "_probe" / "japan.md"
H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0 Safari/537.36"}
s = requests.Session()
parts = ["# Japan probe 2\n"]
for y in (2024, 2023):
    u = f"https://www.esri.cao.go.jp/jp/sna/data/data_list/kenmin/files/contents/main_{y}.html"
    r = s.get(u, headers=H, timeout=60)
    parts.append(f"- {u}: HTTP {r.status_code}")
for name in ("soukatu1", "soukatu9"):
    u = f"https://www.esri.cao.go.jp/jp/sna/data/data_list/kenmin/files/contents/tables/2022/{name}.xlsx"
    wb = openpyxl.load_workbook(io.BytesIO(s.get(u, headers=H, timeout=60).content), read_only=True, data_only=True)
    for ws in wb.worksheets[:2]:
        parts.append(f"\n## {name} sheet {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 80:
                break
            parts.append("- " + " | ".join("" if v is None else str(v)[:14] for v in row[:16]))
body = s.get("https://population.un.org/wup/assets/Download/Cities/WUP2025-F21-DEGURBA-Cities_Pop.xlsx", headers=H, timeout=120).content
ws = openpyxl.load_workbook(io.BytesIO(body), read_only=True)["Data"]
header, rows = None, []
for r in ws.iter_rows(values_only=True):
    if header is None:
        if r and r[0] == "Index":
            header = [str(h) for h in r]
            parts.append(f"\nUN header sample: {header[:14]}")
        continue
    d = dict(zip(header, r))
    if d.get("ISO3_Code") == "JPN" and (d.get("2025") or 0) >= 300:
        rows.append((d["City_Name"], d.get("2025")))
parts.append("## UN WUP 2025 Japanese cities >= 300k")
parts += [f"- {n}: {v:,.0f}" for n, v in sorted(rows, key=lambda t: -t[1])]
OUT.write_text("\n".join(parts), encoding="utf-8")
